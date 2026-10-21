"""Lager ordliste.xlsx fra ordliste.json, et Excel-ark der man kan redigere ordene.

Bruk: python3 lag_excel.py
Etter redigering leser fra_excel.py arket tilbake til ordliste.json og ordliste.csv.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HER = Path(__file__).resolve().parent
RADER = 2000  # rader med nedtrekkslister, så nye ord får dem også
ORDKLASSER = ["substantiv", "verb", "adjektiv", "adverb", "determinativ", "preposisjon", "pronomen", "subjunksjon", "konjunksjon", "interjeksjon"]

# (overskrift i arket, felt i json, bredde, forklaring)
KOLONNER = [
    ("Bokmål", "bokmal", 16, "Bokmålsordet. Spilleren ser det fra bokmål til nynorsk, og det er fasiten fra nynorsk til bokmål."),
    ("Nynorsk", "nynorsk", 16, "Nynorskordet. Spilleren ser det fra nynorsk til bokmål, og det er fasiten fra bokmål til nynorsk."),
    ("Andre riktige nynorskformer", "godkjente_alternativer", 30, "Andre nynorske svar som skal godtas, skilt med semikolon. Eksempel: vera; vere"),
    ("Andre riktige bokmålsformer", "bokmal_alternativer", 30, "Andre bokmålssvar som skal godtas fra nynorsk til bokmål, skilt med semikolon. Eksempel: golv"),
    ("Ordklasse", "ordklasse", 14, "Vises på ordkortet i spillet. Velg fra lista."),
    ("Bokmålsordet godtas på nynorsk", "bokmalsform_godtas", 16, "ja hvis bokmålsordet også er riktig nynorsk med samme betydning (f.eks. tenke). Da godtas det som svar på nynorsk."),
    ("Nynorskordet er lov på bokmål", "nynorskform_lov_pa_bokmal", 16, "ja hvis nynorskordet også er tillatt på bokmål (f.eks. heim). Da godtas det som svar på bokmål."),
    ("Annet ord på nynorsk", "ulikt_ord", 16, "ja hvis bokmålsordet ikke kan oversettes direkte: nynorsk bruker et annet ord eller en annen ending, og bokmålsordet står ikke i Nynorskordboka (f.eks. benytte → nytte, kjærlighet → kjærleik). Ordene får en egen kategori i spillet."),
    ("Lenke til Nynorskordboka", "nynorskordboka", 34, "Fasiten i spillet lenker hit. Står den tom, lager spillet en lenke til oppslaget i Nynorskordboka."),
    ("Merknad", "merknad", 46, "Fritt notat. Vises ikke i spillet."),
]

ARIAL = "Arial"
MORK = "13243C"
INNDATA = PatternFill("solid", fgColor="FFFFFF")


def lag(utfil=HER / "ordliste.xlsx"):
    ord_ = json.loads((HER / "ordliste.json").read_text(encoding="utf-8"))
    wb = Workbook()

    # ---------- Ordliste ----------
    ws = wb.active
    ws.title = "Ordliste"
    for c, (navn, _, bredde, forklaring) in enumerate(KOLONNER, 1):
        cell = ws.cell(row=1, column=c, value=navn)
        cell.font = Font(name=ARIAL, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=MORK)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.comment = Comment(forklaring, "Ordbua")
        ws.column_dimensions[get_column_letter(c)].width = bredde
    ws.row_dimensions[1].height = 34
    for r, o in enumerate(ord_, 2):
        for c, (_, felt, _, _) in enumerate(KOLONNER, 1):
            v = o.get(felt, "")
            if isinstance(v, list):
                v = "; ".join(v)
            cell = ws.cell(row=r, column=c, value=v or None)
            cell.font = Font(name=ARIAL, bold=c <= 2)
            cell.alignment = Alignment(vertical="top", wrap_text=c == len(KOLONNER))
    siste = len(ord_) + 1
    tabell = Table(displayName="Ordliste", ref=f"A1:{get_column_letter(len(KOLONNER))}{siste}")
    tabell.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tabell)
    ws.freeze_panes = "C2"

    dv_klasse = DataValidation(type="list", formula1='"' + ",".join(ORDKLASSER) + '"', allow_blank=False,
                               showErrorMessage=True, errorTitle="Ukjent ordklasse", error="Velg en ordklasse fra lista.")
    dv_jn = DataValidation(type="list", formula1='"ja,nei"', allow_blank=False,
                           showErrorMessage=True, errorTitle="Skriv ja eller nei", error="Denne kolonnen skal være ja eller nei.")
    ws.add_data_validation(dv_klasse)
    ws.add_data_validation(dv_jn)
    dv_klasse.add(f"E2:E{RADER}")
    dv_jn.add(f"F2:H{RADER}")

    # ---------- Slik bruker du arket ----------
    hj = wb.create_sheet("Slik bruker du arket")
    hj.column_dimensions["A"].width = 34
    hj.column_dimensions["B"].width = 90
    tynn = Side(style="thin", color="D3DDE8")
    rad = 1

    def skriv(a, b=None, *, fet=False, storrelse=10, farge=None, fyll=None):
        nonlocal rad
        ca = hj.cell(row=rad, column=1, value=a)
        ca.font = Font(name=ARIAL, bold=True if b is not None or fet else False, size=storrelse, color=farge)
        ca.alignment = Alignment(vertical="top", wrap_text=True)
        if b is not None:
            cb = hj.cell(row=rad, column=2, value=b)
            cb.font = Font(name=ARIAL, size=storrelse, color=farge)
            cb.alignment = Alignment(vertical="top", wrap_text=True)
        if fyll:
            for c in (1, 2):
                hj.cell(row=rad, column=c).fill = PatternFill("solid", fgColor=fyll)
        rad += 1

    skriv("Ordlista til Ordbua", fet=True, storrelse=16)
    skriv("Alle ordene i spillet står i arket «Ordliste». Hver rad er ett ordpar, og spillet bruker det begge veier: "
          "fra bokmål til nynorsk og fra nynorsk til bokmål. Du kan endre, legge til og slette rader der.")
    hj.merge_cells(start_row=rad - 1, start_column=1, end_row=rad - 1, end_column=2)
    hj.row_dimensions[rad - 1].height = 30
    rad += 1

    skriv("Slik gjør du", fet=True, storrelse=12)
    for a, b in [
        ("Endre et ord", "Skriv rett i cellen. Store og små bokstaver spiller ingen rolle i spillet."),
        ("Legge til et ord", "Skriv i første tomme rad under tabellen. Tabellen utvider seg selv. Bokmål, Nynorsk og Ordklasse må fylles ut, og de tre ja/nei-kolonnene får nei hvis de står tomme (de to første får ja hvis ordet er likt på begge)."),
        ("Slette et ord", "Høyreklikk på raden og velg Slett > Tabellrader."),
        ("Flere riktige svar", "Skriv dem i kolonnene «Andre riktige nynorskformer» eller «Andre riktige bokmålsformer», skilt med semikolon (;)."),
        ("Ord som er like", "Skriv samme ord i Bokmål og Nynorsk (f.eks. bil og bil). Spillet merker det som likt på bokmål og nynorsk, og omtrent hvert tredje ord i en runde er et slikt ord."),
        ("Forklaring på kolonnene", "Hold musa over en kolonneoverskrift for å se hva kolonnen betyr."),
        ("Når du er ferdig", "Lagre fila og last den opp i tråden med Claude, så oppdateres spillet. Fila leses tilbake med skriptet fra_excel.py i mappa ordliste."),
    ]:
        skriv(a, b)
        for c in (1, 2):
            hj.cell(row=rad - 1, column=c).border = Border(bottom=tynn)
    rad += 1

    skriv("Eksempel på en rad", fet=True, storrelse=12)
    eksempel = ["hvordan", "korleis", "", "hvorledes", "adverb", "nei", "nei", "ja", "https://ordbokene.no/nno/nn/korleis", ""]
    for (navn, _, _, _), v in zip(KOLONNER, eksempel):
        skriv(navn, v or "(tom)", farge="4B5D75")
    rad += 1

    skriv("Ordlista nå", fet=True, storrelse=12)
    skriv("Ordpar", "=COUNTA(Ordliste!A2:A5000)")
    skriv("Skrives likt på bokmål og nynorsk", '=SUMPRODUCT((Ordliste!A2:A5000<>"")*(LOWER(TRIM(Ordliste!A2:A5000))=LOWER(TRIM(Ordliste!B2:B5000))))')
    skriv("Bokmålsordet godtas på nynorsk", '=COUNTIF(Ordliste!F2:F5000,"ja")')
    skriv("Nynorskordet er lov på bokmål", '=COUNTIF(Ordliste!G2:G5000,"ja")')
    skriv("Annet ord på nynorsk", '=COUNTIF(Ordliste!H2:H5000,"ja")')
    for r in range(rad - 5, rad):
        hj.cell(row=r, column=2).alignment = Alignment(horizontal="left")

    wb.calculation.fullCalcOnLoad = True  # Excel regner ut tallene i «Ordlista nå» når fila åpnes
    wb.save(utfil)
    return utfil, len(ord_)


if __name__ == "__main__":
    f, n = lag()
    print(f"Skrev {f} med {n} ord")
